"""PDF and Word (.docx) export of summarization results.

PDF output needs Unicode TTF fonts to render Indic scripts: the exporter looks
for Noto fonts (Linux: `fonts-noto-core`), Nirmala UI (Windows) or any font
directory set in PDF_FONT_DIR, and switches font per script run. Correct
conjunct shaping additionally needs the `uharfbuzz` package.
"""
import datetime
import io
import logging
import os
import re
import sys
from xml.sax.saxutils import escape

import config

log = logging.getLogger(__name__)

try:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _DOCX_AVAILABLE = True
except Exception:
    docx = None
    _DOCX_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
    _REPORTLAB_AVAILABLE = True
except Exception:
    _REPORTLAB_AVAILABLE = False


class ExportUnavailable(Exception):
    """Raised when the library needed for an export format is not installed."""


# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------

# Unicode block ranges for the scripts the app can output.
_SCRIPT_RANGES = [
    ('devanagari', 0x0900, 0x097F),
    ('bengali', 0x0980, 0x09FF),
    ('gurmukhi', 0x0A00, 0x0A7F),
    ('gujarati', 0x0A80, 0x0AFF),
    ('oriya', 0x0B00, 0x0B7F),
    ('tamil', 0x0B80, 0x0BFF),
    ('telugu', 0x0C00, 0x0C7F),
    ('kannada', 0x0C80, 0x0CFF),
    ('malayalam', 0x0D00, 0x0D7F),
    ('arabic', 0x0600, 0x06FF),
    ('arabic', 0x0750, 0x077F),
    ('arabic', 0xFB50, 0xFDFF),
    ('arabic', 0xFE70, 0xFEFF),
]

# Candidate font files per script, in order of preference (case-insensitive file names).
_FONT_CANDIDATES = {
    'latin': ['NotoSans-Regular.ttf', 'DejaVuSans.ttf', 'arial.ttf', 'Arial.ttf', 'LiberationSans-Regular.ttf'],
    'latin-bold': ['NotoSans-Bold.ttf', 'DejaVuSans-Bold.ttf', 'arialbd.ttf', 'Arial Bold.ttf',
                   'LiberationSans-Bold.ttf'],
    'devanagari': ['NotoSansDevanagari-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'mangal.ttf',
                   'Lohit-Devanagari.ttf'],
    'bengali': ['NotoSansBengali-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'vrinda.ttf', 'Lohit-Bengali.ttf'],
    'gurmukhi': ['NotoSansGurmukhi-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'raavi.ttf',
                 'Lohit-Gurmukhi.ttf'],
    'gujarati': ['NotoSansGujarati-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'shruti.ttf',
                 'Lohit-Gujarati.ttf'],
    'oriya': ['NotoSansOriya-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'kalinga.ttf', 'Lohit-Odia.ttf'],
    'tamil': ['NotoSansTamil-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'latha.ttf', 'Lohit-Tamil.ttf'],
    'telugu': ['NotoSansTelugu-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'gautami.ttf', 'Lohit-Telugu.ttf'],
    'kannada': ['NotoSansKannada-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'tunga.ttf', 'Lohit-Kannada.ttf'],
    'malayalam': ['NotoSansMalayalam-Regular.ttf', 'Nirmala.ttf', 'NirmalaUI.ttf', 'kartika.ttf',
                  'Lohit-Malayalam.ttf'],
    'arabic': ['NotoNaskhArabic-Regular.ttf', 'NotoSansArabic-Regular.ttf', 'arial.ttf', 'Arial.ttf'],
}


def _font_dirs():
    dirs = [config.PDF_FONT_DIR, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fonts')]
    if sys.platform.startswith('win'):
        windir = os.environ.get('WINDIR', r'C:\Windows')
        dirs += [os.path.join(windir, 'Fonts'),
                 os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Microsoft', 'Windows', 'Fonts')]
    elif sys.platform == 'darwin':
        dirs += ['/Library/Fonts', '/System/Library/Fonts', '/System/Library/Fonts/Supplemental',
                 os.path.expanduser('~/Library/Fonts')]
    else:
        dirs += ['/usr/share/fonts', '/usr/local/share/fonts', os.path.expanduser('~/.local/share/fonts'),
                 os.path.expanduser('~/.fonts')]
    return [d for d in dirs if d and os.path.isdir(d)]


_FONT_INDEX = None
_REGISTERED = {}  # script -> registered font name (or None if no font found)


def _font_index():
    global _FONT_INDEX
    if _FONT_INDEX is None:
        _FONT_INDEX = {}
        for base in _font_dirs():
            for root, _dirs, files in os.walk(base):
                for name in files:
                    if name.lower().endswith('.ttf'):
                        _FONT_INDEX.setdefault(name.lower(), os.path.join(root, name))
    return _FONT_INDEX


def _font_for(script):
    """Register and return a font name for script, or None if no font is installed."""
    if script in _REGISTERED:
        return _REGISTERED[script]
    index = _font_index()
    font_name = None
    for candidate in _FONT_CANDIDATES.get(script, []):
        path = index.get(candidate.lower())
        if not path:
            continue
        try:
            font_name = 'RMTC-' + script
            pdfmetrics.registerFont(TTFont(font_name, path))
            # No bold/italic files are loaded for scripts; map <b>/<i> back to regular.
            pdfmetrics.registerFontFamily(font_name, normal=font_name, bold=font_name,
                                          italic=font_name, boldItalic=font_name)
            break
        except Exception as e:
            log.warning("Could not register font %s: %s", path, e)
            font_name = None
    _REGISTERED[script] = font_name
    return font_name


def _base_fonts():
    regular = _font_for('latin') or 'Helvetica'
    bold = _font_for('latin-bold') or ('Helvetica-Bold' if regular == 'Helvetica' else regular)
    if regular != 'Helvetica':
        pdfmetrics.registerFontFamily(regular, normal=regular, bold=bold, italic=regular, boldItalic=bold)
    return regular, bold


def _script_of(ch):
    cp = ord(ch)
    for script, lo, hi in _SCRIPT_RANGES:
        if lo <= cp <= hi:
            return script
    return None


def _markup(text, missing_scripts):
    """Escape text for a reportlab Paragraph, wrapping non-Latin runs in <font> tags.

    Fonts are chosen per whitespace-separated word: reportlab cannot shape a
    word whose characters span two fonts, so punctuation attached to a word
    (e.g. "(ಕನ್ನಡ)") is drawn in that word's script font. Scripts with no
    installed font are recorded in missing_scripts.
    """
    runs = []  # [font or None, text]; None means the paragraph's base font
    for token in re.split(r'(\s+)', text):
        if not token:
            continue
        font = None
        if not token.isspace():
            script = next((s for s in map(_script_of, token) if s), None)
            if script:
                font = _font_for(script)
                if not font:
                    missing_scripts.add(script)
                elif not _covers(font, token):
                    font = None  # mixed-script word; base font avoids missing glyphs for Latin
        if runs and (runs[-1][0] == font or token.isspace()):
            runs[-1][1] += token
        else:
            runs.append([font, token])

    parts = []
    for font, chunk in runs:
        safe = escape(chunk).replace('\n', '<br/>')
        parts.append(f'<font name="{font}">{safe}</font>' if font else safe)
    return ''.join(parts)


def _covers(font_name, token):
    cmap = pdfmetrics.getFont(font_name).face.charToGlyph
    return all(ord(ch) in cmap for ch in token)


def _is_rtl(text):
    letters = [ch for ch in text if ch.isalpha()]
    return bool(letters) and sum(1 for ch in letters if _script_of(ch) == 'arabic') > len(letters) / 2


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def generate_docx_bytes(summary, original_text="", language_name="English", keywords=None, explanation=""):
    """Return a Word (.docx) document containing the summary output.

    Raises ExportUnavailable when python-docx is not installed.
    """
    if not _DOCX_AVAILABLE:
        raise ExportUnavailable("Word export needs the python-docx package.")

    doc = docx.Document()

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Text Summarization & Translation Output")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    time_str = datetime.datetime.now().strftime("%B %d, %Y - %H:%M")
    meta_run = meta_p.add_run(f"Generated on: {time_str} | Language: {language_name}")
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    h1 = doc.add_heading("1. Summary Result", level=1)
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    summary_run = doc.add_paragraph().add_run(summary or "No summary generated.")
    summary_run.font.size = Pt(11)
    summary_run.font.color.rgb = RGBColor(15, 23, 42)

    if keywords:
        h2 = doc.add_heading("2. Extracted Keywords", level=1)
        h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        kw_run = doc.add_paragraph().add_run(", ".join(str(k) for k in keywords))
        kw_run.font.size = Pt(10.5)
        kw_run.font.bold = True
        kw_run.font.color.rgb = RGBColor(79, 70, 229)

    if explanation:
        h3 = doc.add_heading("3. Detailed Explanation", level=1)
        h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().add_run(explanation).font.size = Pt(10.5)

    if original_text:
        doc.add_page_break()
        h4 = doc.add_heading("Original Input Text", level=1)
        h4.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        orig_run = doc.add_paragraph().add_run(
            original_text[:4000] + ("..." if len(original_text) > 4000 else ""))
        orig_run.font.size = Pt(9.5)
        orig_run.font.color.rgb = RGBColor(100, 116, 139)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_pdf_bytes(summary, original_text="", language_name="English", keywords=None, explanation=""):
    """Return a PDF document containing the summary output.

    Raises ExportUnavailable when reportlab is not installed. Text in scripts
    with no installed font is logged; it renders as blank boxes.
    """
    if not _REPORTLAB_AVAILABLE:
        raise ExportUnavailable("PDF export needs the reportlab package.")

    regular, bold = _base_fonts()
    missing = set()
    styles = getSampleStyleSheet()

    def style(name, parent, **kw):
        kw.setdefault('fontName', regular)
        return ParagraphStyle(name, parent=styles[parent], shaping=1, **kw)

    title_style = style('DocTitle', 'Heading1', fontName=bold, fontSize=20, leading=24,
                        textColor=colors.HexColor('#1e293b'), alignment=1, spaceAfter=6)
    meta_style = style('DocMeta', 'Normal', fontSize=10, leading=14,
                       textColor=colors.HexColor('#64748b'), alignment=1, spaceAfter=15)
    heading_style = style('DocHeading', 'Heading2', fontName=bold, fontSize=14, leading=18,
                          textColor=colors.HexColor('#4f46e5'), spaceBefore=12, spaceAfter=6)
    body_style = style('DocBody', 'Normal', fontSize=10.5, leading=17,
                       textColor=colors.HexColor('#0f172a'), spaceAfter=10)
    kw_style = style('DocKeywords', 'Normal', fontSize=10, leading=15,
                     textColor=colors.HexColor('#4338ca'), spaceAfter=10)
    orig_style = style('OrigText', 'Normal', fontSize=9, leading=14,
                       textColor=colors.HexColor('#475569'), spaceAfter=10)

    def para(text, st):
        if _is_rtl(text):
            st = ParagraphStyle(st.name + 'RTL', parent=st, alignment=2)
        return Paragraph(_markup(text, missing), st)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40,
                            title="Text Summarization Output")

    time_str = datetime.datetime.now().strftime("%B %d, %Y - %H:%M")
    elements = [
        Paragraph("Text Summarization &amp; Translation Output", title_style),
        Paragraph(f"Language: {_markup(str(language_name), missing)} | Generated on: {time_str}",
                  meta_style),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0'), spaceAfter=15),
        Paragraph("Summary Output", heading_style),
        para(summary or "No summary generated.", body_style),
        Spacer(1, 10),
    ]

    if keywords:
        elements += [
            Paragraph("Extracted Keywords", heading_style),
            para(", ".join(str(k) for k in keywords), kw_style),
            Spacer(1, 10),
        ]

    if explanation:
        elements += [
            Paragraph("Summary Explanation", heading_style),
            para(explanation, body_style),
            Spacer(1, 10),
        ]

    if original_text:
        excerpt = original_text[:3000] + ("..." if len(original_text) > 3000 else "")
        elements += [
            Spacer(1, 15),
            HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cbd5e1'), spaceAfter=10),
            Paragraph("Original Input Source", heading_style),
            para(excerpt, orig_style),
        ]

    if missing:
        log.warning("No font installed for script(s) %s; install Noto fonts or set PDF_FONT_DIR.",
                    ', '.join(sorted(missing)))

    doc.build(elements)
    return buffer.getvalue()
