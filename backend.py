"""Hybrid extractive + abstractive summarization with multilingual output.

Pipeline (see docs/ARCHITECTURE.md for details):
    detect language -> translate to English -> split sentences -> embed ->
    select salient, diverse sentences (MMR) -> rewrite with BART ->
    keywords + explanation -> translate to target language

Every heavy dependency is optional. Without torch/transformers (or with
LIGHTWEIGHT_MODE=1) the pipeline falls back to TF-IDF selection and
frequency-based extraction, and to Google Translate for translation.
"""
import collections
import io
import logging
import os
import re
import sys
import threading
import time
import warnings
from dataclasses import dataclass, field

import config

log = logging.getLogger(__name__)

# Ensure UTF-8 output encoding for Windows terminals handling non-ASCII translations
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

os.environ.setdefault("TF_ENABLE_ONEDNN_OPTS", "0")

# Optional heavy ML libs: import where available, otherwise set flags for fallback behavior
try:
    import torch
    _TORCH_AVAILABLE = True
except Exception:
    torch = None
    _TORCH_AVAILABLE = False

try:
    from transformers import (
        AutoModelForSeq2SeqLM,
        AutoTokenizer,
        BartForConditionalGeneration,
        BartTokenizer,
        MarianMTModel,
        MarianTokenizer,
    )
    _HF_AVAILABLE = True
except Exception:
    AutoModelForSeq2SeqLM = AutoTokenizer = None
    BartForConditionalGeneration = BartTokenizer = None
    MarianMTModel = MarianTokenizer = None
    _HF_AVAILABLE = False

try:
    from sentence_transformers import SentenceTransformer
    _SENTENCE_TRANSFORMERS_AVAILABLE = True
except Exception:
    SentenceTransformer = None
    _SENTENCE_TRANSFORMERS_AVAILABLE = False

try:
    import numpy as np
except Exception:
    np = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
except Exception:
    TfidfVectorizer = None

try:
    from nltk.tokenize import sent_tokenize
except Exception:
    sent_tokenize = None

try:
    from nltk.corpus import stopwords as _nltk_stopwords
    _NLTK_STOPWORDS = set(_nltk_stopwords.words('english'))
except Exception:
    _NLTK_STOPWORDS = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    import docx
    _DOCX_AVAILABLE = True
except Exception:
    docx = None
    _DOCX_AVAILABLE = False

try:
    from langdetect import DetectorFactory, detect_langs
    DetectorFactory.seed = 0  # langdetect is non-deterministic without a fixed seed
    _LANGDETECT_AVAILABLE = True
except Exception:
    _LANGDETECT_AVAILABLE = False

try:
    from PIL import Image
    import pytesseract
    OCR_AVAILABLE = True
except Exception:
    OCR_AVAILABLE = False

try:
    from deep_translator import GoogleTranslator
    from deep_translator.exceptions import TooManyRequests
except Exception:
    GoogleTranslator = None
    TooManyRequests = None

if _TORCH_AVAILABLE:
    DEVICE = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
else:
    DEVICE = None

SUMMARY_MODES = ('simple', 'balanced', 'detailed')
_MODE_ALIASES = {
    'easy': 'simple', 'easy_understand': 'simple', 'short': 'simple',
    'long': 'detailed', 'verbose': 'detailed', 'technical': 'detailed',
}

# Model names (overridable via environment, see .env.example)
SUMMARIZER_MODEL = os.environ.get('SUMMARIZER_MODEL', 'facebook/bart-large-cnn')
EMBEDDING_MODEL = os.environ.get('EMBEDDING_MODEL', 'paraphrase-multilingual-MiniLM-L12-v2')
MULTILINGUAL_MODEL = os.environ.get('MULTILINGUAL_MODEL', 'facebook/nllb-200-distilled-600M')
TO_ENGLISH_MODEL = 'Helsinki-NLP/opus-mt-mul-en'

# English -> target Marian models: (model id, required target-language prefix token).
# Only models that actually exist on the Hugging Face Hub are listed.
_MARIAN_FROM_EN = {
    'hi': ('Helsinki-NLP/opus-mt-en-hi', None),
    'ml': ('Helsinki-NLP/opus-mt-en-ml', None),
    'mr': ('Helsinki-NLP/opus-mt-en-mr', None),
    'ur': ('Helsinki-NLP/opus-mt-en-ur', None),
    'es': ('Helsinki-NLP/opus-mt-en-es', None),
    'fr': ('Helsinki-NLP/opus-mt-en-fr', None),
    'de': ('Helsinki-NLP/opus-mt-en-de', None),
    'kn': ('Helsinki-NLP/opus-mt-en-dra', '>>kan<<'),
    'ta': ('Helsinki-NLP/opus-mt-en-dra', '>>tam<<'),
    'te': ('Helsinki-NLP/opus-mt-en-dra', '>>tel<<'),
    'bn': ('Helsinki-NLP/opus-mt-en-inc', '>>ben<<'),
    'gu': ('Helsinki-NLP/opus-mt-en-inc', '>>guj<<'),
    'pa': ('Helsinki-NLP/opus-mt-en-inc', '>>pan_Guru<<'),
    'or': ('Helsinki-NLP/opus-mt-en-inc', '>>ori<<'),
    'ne': ('Helsinki-NLP/opus-mt-en-inc', '>>npi<<'),
}

# Kannada, Telugu and Tamil: opus-mt-mul-en invents unrelated English for them
# and mBART-50 lacks Kannada and mixes scripts in Telugu, so these languages go
# NLLB -> Google -> Marian (English -> language only).
_NLLB_FIRST = {'kn', 'ta', 'te'}

# Unicode blocks used to recognise these scripts before statistical detection.
_SCRIPT_RANGES = {
    'kn': ('\u0c80', '\u0cff'),
    'te': ('\u0c00', '\u0c7f'),
    'ta': ('\u0b80', '\u0bff'),
}

# Language codes for the multilingual model.
# mBART-50 has no Kannada, Punjabi or Odia; NLLB-200 covers all of them.
_MBART50_CODES = {
    'en': 'en_XX', 'hi': 'hi_IN', 'bn': 'bn_IN', 'gu': 'gu_IN', 'ml': 'ml_IN',
    'mr': 'mr_IN', 'ta': 'ta_IN', 'te': 'te_IN', 'ur': 'ur_PK', 'ne': 'ne_NP',
    'es': 'es_XX', 'fr': 'fr_XX', 'de': 'de_DE',
}
_NLLB_CODES = {
    'en': 'eng_Latn', 'hi': 'hin_Deva', 'bn': 'ben_Beng', 'gu': 'guj_Gujr',
    'kn': 'kan_Knda', 'ml': 'mal_Mlym', 'mr': 'mar_Deva', 'or': 'ory_Orya',
    'pa': 'pan_Guru', 'ta': 'tam_Taml', 'te': 'tel_Telu', 'ur': 'urd_Arab',
    'ne': 'npi_Deva', 'es': 'spa_Latn', 'fr': 'fra_Latn', 'de': 'deu_Latn',
}

# Text the translators must not touch: URLs, emails, @handles, phone numbers.
# They are swapped for [[n]] placeholders, which NLLB copies through unchanged.
_PROTECTED = re.compile(
    r'(?:https?://|www\.)\S+?(?=[.,;:!?)\]]*(?:\s|$))'
    r'|[\w.+-]+@[\w-]+(?:\.[\w-]+)+'
    r'|(?<![\w@])@\w+'
    r'|\+?\d[\d -]{7,}\d')

# Non-English documents with more sentences than this are cut down to their most
# central ones (TF-IDF in the source language) before translation to English,
# which otherwise costs 1-3 s per sentence on CPU.
_PRESELECT_SENTENCES = 40

# Token budgets for model-based translation chunks (Marian max input is 512).
_MT_CHUNK_TOKENS = 400
_MT_BATCH_SIZE = 8
_ONLINE_CHUNK_CHARS = 4500  # Google Translate rejects requests over 5000 chars

_FALLBACK_STOPWORDS = {
    'a', 'about', 'above', 'after', 'again', 'against', 'all', 'also', 'am', 'an', 'and', 'any',
    'are', 'as', 'at', 'be', 'because', 'been', 'before', 'being', 'below', 'between', 'both',
    'but', 'by', 'can', 'could', 'did', 'do', 'does', 'doing', 'down', 'during', 'each', 'few',
    'for', 'from', 'further', 'had', 'has', 'have', 'having', 'he', 'her', 'here', 'hers',
    'him', 'his', 'how', 'however', 'into', 'is', 'it', 'its', 'itself', 'just', 'many', 'may',
    'more', 'most', 'much', 'must', 'not', 'now', 'off', 'once', 'only', 'other', 'our',
    'ours', 'out', 'over', 'own', 'same', 'she', 'should', 'some', 'such', 'than', 'that',
    'the', 'their', 'theirs', 'them', 'then', 'there', 'these', 'they', 'this', 'those',
    'through', 'too', 'under', 'until', 'upon', 'very', 'was', 'were', 'what', 'when', 'where',
    'which', 'while', 'who', 'whom', 'why', 'will', 'with', 'within', 'without', 'would',
    'you', 'your', 'yours', 'one', 'two', 'new', 'used', 'use', 'using', 'well', 'like',
    'make', 'made', 'get', 'said', 'says', 'even', 'still', 'yet', 'via', 'per',
}
STOPWORDS = _NLTK_STOPWORDS | _FALLBACK_STOPWORDS if _NLTK_STOPWORDS else _FALLBACK_STOPWORDS

# Model caches. Loading is guarded by a lock so concurrent first requests do not
# load the same multi-GB model twice; translation with a shared tokenizer whose
# src_lang must be set per call is guarded by its own lock.
_LOAD_LOCK = threading.RLock()
_MULTI_LOCK = threading.Lock()
_EMBED_MODEL = None
_EMBED_LOADED = False
_ABSTR = None           # (model, tokenizer) or False once loading failed
_MARIAN = {}            # model id -> (model, tokenizer) or None once loading failed
_MULTI = None           # (model, tokenizer) or False once loading failed

# Sentence-level translation cache: (src, tgt, chunk) -> translation. Saves
# re-translating the summary inside the explanation, and repeat requests.
_TRANSLATION_CACHE = collections.OrderedDict()
_CACHE_LOCK = threading.Lock()
_CACHE_SIZE = 5000


class ExtractionError(Exception):
    """Raised when no usable text can be extracted from an uploaded file."""


@dataclass
class SummaryResult:
    summary: str
    original_text: str
    explanation: str = ''
    keywords: list = field(default_factory=list)
    keywords_language: str = 'en'
    source_language: str = 'en'
    language: str = 'en'
    method: str = 'extractive'
    warnings: list = field(default_factory=list)

    def to_dict(self):
        return dict(self.__dict__)


def _models_enabled():
    return _HF_AVAILABLE and _TORCH_AVAILABLE and not config.LIGHTWEIGHT_MODE


def _to_device(encoded):
    if DEVICE is None:
        return encoded
    return {k: v.to(DEVICE) for k, v in encoded.items()}


# ---------------------------------------------------------------------------
# Model loading
# ---------------------------------------------------------------------------

def _load_embedding_model():
    global _EMBED_MODEL, _EMBED_LOADED
    if _EMBED_LOADED:
        return _EMBED_MODEL
    with _LOAD_LOCK:
        if not _EMBED_LOADED:
            if _SENTENCE_TRANSFORMERS_AVAILABLE and not config.LIGHTWEIGHT_MODE:
                try:
                    _EMBED_MODEL = SentenceTransformer(EMBEDDING_MODEL)
                except Exception as e:
                    log.warning("Embedding model load failed, using TF-IDF: %s", e)
            _EMBED_LOADED = True
    return _EMBED_MODEL


def _load_abstractive_model():
    global _ABSTR
    if _ABSTR is not None:
        return _ABSTR or None
    if not _models_enabled():
        return None
    with _LOAD_LOCK:
        if _ABSTR is None:
            try:
                tokenizer = BartTokenizer.from_pretrained(SUMMARIZER_MODEL)
                model = BartForConditionalGeneration.from_pretrained(SUMMARIZER_MODEL)
                if DEVICE is not None:
                    model.to(DEVICE)
                model.eval()
                _ABSTR = (model, tokenizer)
            except Exception as e:
                log.warning("Summarizer model load failed, using extractive output: %s", e)
                _ABSTR = False
    return _ABSTR or None


def _load_marian(model_id):
    if model_id in _MARIAN:
        return _MARIAN[model_id]
    if not _models_enabled():
        return None
    with _LOAD_LOCK:
        if model_id not in _MARIAN:
            try:
                tokenizer = MarianTokenizer.from_pretrained(model_id)
                model = MarianMTModel.from_pretrained(model_id)
                if DEVICE is not None:
                    model.to(DEVICE)
                model.eval()
                _MARIAN[model_id] = (model, tokenizer)
            except Exception as e:
                log.warning("Translation model %s failed to load: %s", model_id, e)
                _MARIAN[model_id] = None
    return _MARIAN[model_id]


def _load_multilingual():
    global _MULTI
    if _MULTI is not None:
        return _MULTI or None
    if not _models_enabled():
        return None
    with _LOAD_LOCK:
        if _MULTI is None:
            try:
                tokenizer = AutoTokenizer.from_pretrained(MULTILINGUAL_MODEL)
                model = AutoModelForSeq2SeqLM.from_pretrained(MULTILINGUAL_MODEL)
                if DEVICE is not None:
                    model.to(DEVICE)
                model.eval()
                if config.QUANTIZE_TRANSLATION and DEVICE is not None and DEVICE.type == 'cpu':
                    model = _quantize_int8(model)
                _MULTI = (model, tokenizer)
            except Exception as e:
                log.warning("Multilingual model %s failed to load: %s", MULTILINGUAL_MODEL, e)
                _MULTI = False
    return _MULTI or None


def _quantize_int8(model):
    """int8 dynamic quantization of Linear layers: about 2x faster on CPU.

    Only used for the translation model; BART produces nonsense when quantized.
    """
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')  # torch deprecation notice for the eager API
        return torch.ao.quantization.quantize_dynamic(model, {torch.nn.Linear}, dtype=torch.qint8)


def preload_models():
    """Load every model the pipeline uses, so the first request is not slow."""
    start = time.perf_counter()
    _load_embedding_model()
    _load_abstractive_model()
    _load_multilingual()
    log.info("Models ready after %.0f s", time.perf_counter() - start)


def _multilingual_is_nllb():
    return 'nllb' in MULTILINGUAL_MODEL.lower()


def _multilingual_codes():
    return _NLLB_CODES if _multilingual_is_nllb() else _MBART50_CODES


# ---------------------------------------------------------------------------
# Translation
# ---------------------------------------------------------------------------

def _is_break(chunk):
    return not chunk.strip() and '\n' in chunk


def _chunk_text(text, measure, limit, per_sentence=False):
    """Group sentences into chunks whose measure() stays under limit.

    Over-long sentences are split on whitespace. With per_sentence=True every
    sentence becomes its own chunk. Line breaks are returned as separate chunks
    (e.g. '\n' or '\n\n') so translated output keeps the original layout;
    callers must pass them through untranslated.
    """
    chunks = []
    for block in re.split(r'(\n\s*)', text):
        if _is_break(block):
            if chunks:
                chunks.append('\n\n' if block.count('\n') > 1 else '\n')
            continue
        block = block.strip()
        if not block:
            continue
        current = ''
        for sentence in _split_sentences(block):
            pieces = [sentence]
            if measure(sentence) > limit:
                pieces, part = [], ''
                for word in sentence.split():
                    candidate = f"{part} {word}".strip()
                    if part and measure(candidate) > limit:
                        pieces.append(part)
                        part = word
                    else:
                        part = candidate
                if part:
                    pieces.append(part)
            for piece in pieces:
                candidate = f"{current} {piece}".strip()
                if current and (per_sentence or measure(candidate) > limit):
                    chunks.append(current)
                    current = piece
                else:
                    current = candidate
        if current:
            chunks.append(current)
    while chunks and _is_break(chunks[-1]):
        chunks.pop()
    return chunks


def _join_chunks(chunks):
    out = ''
    for c in chunks:
        if _is_break(c):
            out = out.rstrip() + c
        else:
            out += ('' if not out or out.endswith('\n') else ' ') + c
    return out.strip()


def _cache_get(key):
    with _CACHE_LOCK:
        value = _TRANSLATION_CACHE.get(key)
        if value is not None:
            _TRANSLATION_CACHE.move_to_end(key)
        return value


def _cache_put(key, value):
    with _CACHE_LOCK:
        _TRANSLATION_CACHE[key] = value
        _TRANSLATION_CACHE.move_to_end(key)
        while len(_TRANSLATION_CACHE) > _CACHE_SIZE:
            _TRANSLATION_CACHE.popitem(last=False)


def _translate_chunks(text, translate_batch, measure, limit, per_sentence=False, batch_size=1,
                      cache_key=None):
    """Translate text chunk-by-chunk. Returns None if any chunk fails.

    translate_batch receives a list of up to batch_size chunks and returns a
    list of translations in the same order. With cache_key (src, tgt), chunks
    already translated are reused and each distinct chunk is translated once.
    """
    chunks = _chunk_text(text, measure, limit, per_sentence=per_sentence)
    done = {}
    for c in chunks:
        if not _is_break(c) and cache_key:
            hit = _cache_get(cache_key + (c,))
            if hit:
                done[c] = hit
    # Similar lengths in a batch mean less padding.
    pending = sorted({c for c in chunks if not _is_break(c) and c not in done}, key=len)
    for start in range(0, len(pending), batch_size):
        batch = pending[start:start + batch_size]
        translated = translate_batch(batch)
        if not translated or len(translated) != len(batch) or not all(translated):
            return None
        for c, t in zip(batch, translated):
            done[c] = t
            if cache_key:
                _cache_put(cache_key + (c,), t)
    return _join_chunks([c if _is_break(c) else done[c] for c in chunks])


def _generate(model, tokenizer, texts, **gen_kwargs):
    """Run a seq2seq model on a batch of texts; returns decoded strings."""
    encoded = _to_device(tokenizer(texts, return_tensors='pt', truncation=True, max_length=512, padding=True))
    n_in = encoded['input_ids'].shape[-1]
    with torch.no_grad():
        out = model.generate(**encoded, max_length=min(1024, 3 * n_in + 32), **gen_kwargs)
    return [t.strip() for t in tokenizer.batch_decode(out, skip_special_tokens=True)]


def _translate_marian(text, src_lang, tgt_lang):
    if tgt_lang == 'en':
        if src_lang in _NLLB_FIRST:
            return None
        model_id, prefix = TO_ENGLISH_MODEL, None
    elif src_lang == 'en' and tgt_lang in _MARIAN_FROM_EN:
        model_id, prefix = _MARIAN_FROM_EN[tgt_lang]
    else:
        return None
    loaded = _load_marian(model_id)
    if not loaded:
        return None
    model, tokenizer = loaded

    def batch(chunks):
        return _generate(model, tokenizer, [f"{prefix} {c}" if prefix else c for c in chunks])

    try:
        # One sentence per input: Marian is trained on sentence pairs and drops
        # content when given several sentences at once.
        return _translate_chunks(text, batch, lambda s: len(tokenizer.tokenize(s)), _MT_CHUNK_TOKENS,
                                 per_sentence=True, batch_size=_MT_BATCH_SIZE,
                                 cache_key=(model_id, src_lang, tgt_lang))
    except Exception as e:
        log.warning("Marian %s->%s failed: %s", src_lang, tgt_lang, e)
        return None


def _translate_multilingual(text, src_lang, tgt_lang):
    codes = _multilingual_codes()
    src_tag, tgt_tag = codes.get(src_lang), codes.get(tgt_lang)
    if not src_tag or not tgt_tag:
        return None
    loaded = _load_multilingual()
    if not loaded:
        return None
    model, tokenizer = loaded
    forced_id = tokenizer.convert_tokens_to_ids(tgt_tag)
    if forced_id is None or forced_id == tokenizer.unk_token_id:
        return None

    def batch(chunks):
        # src_lang lives on the shared tokenizer; hold the lock across set + encode.
        with _MULTI_LOCK:
            tokenizer.src_lang = src_tag
            return _generate(model, tokenizer, chunks, forced_bos_token_id=forced_id,
                             num_beams=config.TRANSLATION_BEAMS)

    try:
        return _translate_chunks(text, batch, lambda s: len(tokenizer.tokenize(s)), _MT_CHUNK_TOKENS,
                                 per_sentence=True, batch_size=_MT_BATCH_SIZE,
                                 cache_key=(MULTILINGUAL_MODEL, src_lang, tgt_lang))
    except Exception as e:
        log.warning("%s %s->%s failed: %s", MULTILINGUAL_MODEL, src_lang, tgt_lang, e)
        return None


def _translate_online(text, src_lang, tgt_lang):
    if GoogleTranslator is None or not config.ENABLE_ONLINE_TRANSLATION:
        return None
    try:
        translator = GoogleTranslator(source=src_lang or 'auto', target=tgt_lang)
    except Exception:
        translator = GoogleTranslator(source='auto', target=tgt_lang)

    def translate(chunk):
        try:
            return translator.translate(chunk)
        except Exception as e:
            if TooManyRequests is None or not isinstance(e, TooManyRequests):
                raise
            time.sleep(2)  # Google allows ~5 requests/second; retry once
            return translator.translate(chunk)

    try:
        return _translate_chunks(text, lambda chunks: [translate(c) for c in chunks],
                                 len, _ONLINE_CHUNK_CHARS, cache_key=('google', src_lang, tgt_lang))
    except Exception as e:
        log.warning("Online translation %s->%s failed: %s", src_lang, tgt_lang, e)
        return None


def _translation_backends(src_lang, tgt_lang):
    if src_lang in _NLLB_FIRST or tgt_lang in _NLLB_FIRST:
        if _multilingual_is_nllb():
            return (_translate_multilingual, _translate_online, _translate_marian)
        return (_translate_online, _translate_marian)
    return (_translate_marian, _translate_multilingual, _translate_online)


def _protect(text):
    """Replace URLs, emails, handles and phone numbers with [[n]] placeholders."""
    found = []

    def placeholder(m):
        found.append(m.group(0))
        return f"[[{len(found) - 1}]]"
    return _PROTECTED.sub(placeholder, text), found


def _restore(text, found):
    """Undo _protect; None if the translator dropped or duplicated a placeholder."""
    for i, value in enumerate(found):
        tag = f"[[{i}]]"
        if text.count(tag) != 1:
            return None
        text = text.replace(tag, value)
    return text


def _translate_text(text, src_lang, tgt_lang):
    """Translate text with the best available backend for the language pair.

    Usually Marian, then the multilingual model, then Google; see _NLLB_FIRST
    for the exceptions. Returns the translation, the original text when
    src == tgt, or None when every backend failed.
    """
    if not text or not text.strip() or src_lang == tgt_lang:
        return text
    masked, protected = _protect(text)
    for backend in _translation_backends(src_lang, tgt_lang):
        result = backend(masked, src_lang, tgt_lang)
        if result and protected:
            # Placeholder lost: translate unmasked (unchanged sentences hit the cache).
            result = _restore(result, protected) or backend(text, src_lang, tgt_lang)
        if result:
            return result
    return None


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text(f):
    """Extract text from an uploaded file (werkzeug FileStorage) or a path.

    Raises ExtractionError with a user-facing message when nothing usable is found.
    """
    if isinstance(f, str):
        name = f
        with open(f, 'rb') as fh:
            content = fh.read()
    else:
        name = getattr(f, 'filename', '') or 'uploaded_file'
        content = f.read()
        if hasattr(f, 'seek'):
            f.seek(0)

    if not content:
        raise ExtractionError(f"{name} is empty.")

    lower = name.lower()
    try:
        if lower.endswith('.pdf'):
            if PdfReader is None:
                raise ExtractionError("PDF support is not installed on the server (pypdf).")
            reader = PdfReader(io.BytesIO(content))
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
            if not text.strip():
                raise ExtractionError(
                    f"{name} has no extractable text (it may be a scanned PDF). "
                    "Convert the pages to images and upload them for OCR.")
        elif lower.endswith('.docx'):
            if not _DOCX_AVAILABLE:
                raise ExtractionError("Word document support is not installed on the server (python-docx).")
            document = docx.Document(io.BytesIO(content))
            text = '\n'.join(p.text for p in document.paragraphs if p.text.strip())
        elif lower.endswith(('.jpg', '.jpeg', '.png')):
            if not OCR_AVAILABLE:
                raise ExtractionError("Image OCR is not available on the server (Pillow, pytesseract, tesseract).")
            text = pytesseract.image_to_string(Image.open(io.BytesIO(content)))
        else:
            text = content.decode('utf-8', errors='ignore')
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Could not read {name}: {e}") from e

    if not text.strip():
        raise ExtractionError(f"No text found in {name}.")
    return text


# ---------------------------------------------------------------------------
# Summarization helpers
# ---------------------------------------------------------------------------

def _split_sentences(text):
    if sent_tokenize is not None:
        try:
            return [s.strip() for s in sent_tokenize(text) if s.strip()]
        except LookupError:
            pass  # NLTK punkt data not downloaded
        except Exception:
            pass
    # Regex fallback; also splits on the Devanagari danda used in Hindi/Marathi.
    return [s.strip() for s in re.split(r'(?<=[.!?।])\s+', text) if s.strip()]


def _compute_embeddings(sentences):
    """Return an (n, d) array of L2-normalised sentence vectors, or None."""
    if np is None:
        return None
    vectors = None
    model = _load_embedding_model()
    if model is not None:
        try:
            vectors = model.encode(sentences, convert_to_numpy=True)
        except Exception as e:
            log.warning("Embedding failed, using TF-IDF: %s", e)
    if vectors is None and TfidfVectorizer is not None:
        try:
            vectors = TfidfVectorizer(max_features=2048).fit_transform(sentences).toarray()
        except ValueError:
            return None  # e.g. every sentence was empty after tokenization
    if vectors is None:
        return None
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return vectors / norms


def _hybrid_selector(embeddings, top_k, diversity_lambda=0.6):
    """Maximal Marginal Relevance selection.

    Relevance is cosine similarity to the document centroid; each step
    penalises similarity to sentences already selected. Deterministic.
    Returns (sorted indices, {index: relevance}).
    """
    n = embeddings.shape[0]
    centroid = embeddings.mean(axis=0)
    norm = np.linalg.norm(centroid)
    relevance = embeddings @ (centroid / norm) if norm else np.zeros(n)
    if n <= top_k:
        return list(range(n)), {i: float(relevance[i]) for i in range(n)}

    similarity = embeddings @ embeddings.T
    selected = []
    max_sim = np.zeros(n)
    available = np.ones(n, dtype=bool)
    for _ in range(top_k):
        scores = relevance - (diversity_lambda * max_sim if selected else 0)
        scores = np.where(available, scores, -np.inf)
        best = int(np.argmax(scores))
        selected.append(best)
        available[best] = False
        max_sim = np.maximum(max_sim, similarity[best])
    return sorted(selected), {i: float(relevance[i]) for i in selected}


def _frequency_selector(sentences, top_k):
    """Score sentences by average word frequency; used when vectors are unavailable."""
    freq = collections.Counter(
        w for w in re.findall(r"\w+", ' '.join(sentences).lower()) if w not in STOPWORDS)
    scores = {}
    for i, s in enumerate(sentences):
        words = [w for w in re.findall(r"\w+", s.lower()) if w not in STOPWORDS]
        if not words:
            scores[i] = 0.0
            continue
        scores[i] = sum(freq[w] for w in words) / len(words) * min(1.0, len(words) / 30)
    chosen = sorted(sorted(scores, key=scores.get, reverse=True)[:top_k])
    top = max(scores.values()) or 1.0
    return chosen, {i: scores[i] / top for i in chosen}


def _preselect(sentences, limit):
    """Indices of up to `limit` central, diverse sentences, chosen without translating.

    TF-IDF on whitespace tokens works for any script, unlike the embedding
    model, which was not trained on Kannada, Telugu or Tamil.
    """
    if len(sentences) <= limit or TfidfVectorizer is None or np is None:
        return list(range(len(sentences)))
    try:
        vectors = TfidfVectorizer(token_pattern=r"[^\s\d.,;:!?()\[\]\"'\u0964]+",
                                  max_features=4096).fit_transform(sentences).toarray()
    except ValueError:
        return list(range(len(sentences)))
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return _hybrid_selector(vectors / norms, limit)[0]


def _extract_keywords(text, top_n=10):
    words = [w.lower() for w in re.findall(r"[^\W\d_]{3,}", text)]
    filtered = [w for w in words if w not in STOPWORDS]
    return [w for w, _ in collections.Counter(filtered).most_common(top_n)]


def _detect_language(text):
    """Best-effort ISO 639-1 code for text; defaults to English when unsure."""
    sample = text.strip()[:5000]
    letters = re.findall(r"[^\W\d_]", sample)
    if not letters:
        return 'en'
    # Script-based detection is exact; langdetect can fall below its confidence
    # threshold on Kannada/Telugu/Tamil text that mixes in English words.
    for lang, (lo, hi) in _SCRIPT_RANGES.items():
        if sum(1 for ch in letters if lo <= ch <= hi) / len(letters) > 0.3:
            return lang
    ascii_ratio = sum(1 for ch in letters if ch.isascii()) / len(letters)
    # Short Latin-script snippets are routinely misdetected (e.g. as 'af' or 'so').
    if ascii_ratio > 0.9 and len(sample) < 200:
        return 'en'
    if not _LANGDETECT_AVAILABLE:
        return 'en'
    try:
        best = detect_langs(sample)[0]
    except Exception:
        return 'en'
    if best.prob < 0.85:
        return 'en'
    return best.lang.split('-')[0]


def normalize_mode(mode):
    mode = (mode or config.DEFAULT_SUMMARY_MODE or 'balanced').strip().lower()
    mode = _MODE_ALIASES.get(mode, mode)
    return mode if mode in SUMMARY_MODES else 'balanced'


def _mode_settings(mode, n_sentences):
    if mode == 'simple':
        return dict(top_k=max(2, min(6, n_sentences // 8)), min_length=20, max_length=120,
                    num_beams=2, length_penalty=2.0)
    if mode == 'detailed':
        return dict(top_k=max(6, min(12, n_sentences // 4)), min_length=120, max_length=400,
                    num_beams=6, length_penalty=0.8)
    return dict(top_k=max(3, min(8, n_sentences // 6)), min_length=60, max_length=300,
                num_beams=4, length_penalty=1.0)


def _generation_lengths(n_input_tokens, settings):
    """Return (min_length, max_length) for BART, or None to skip rewriting.

    The mode presets are upper bounds: forcing BART to write more tokens than
    the input contains makes it repeat itself or invent content.
    """
    if n_input_tokens < 40:
        return None  # too short to be worth rewriting
    max_length = max(20, min(settings['max_length'], n_input_tokens))
    min_length = min(settings['min_length'], int(max_length * 0.6))
    return min_length, max_length


def _abstractive_rewrite(text, settings):
    """Rewrite extracted sentences with BART. Returns None when unavailable."""
    loaded = _load_abstractive_model()
    if not loaded:
        return None
    model, tokenizer = loaded
    inputs = _to_device(tokenizer([text], max_length=1024, return_tensors='pt', truncation=True))
    lengths = _generation_lengths(inputs['input_ids'].shape[-1], settings)
    if lengths is None:
        return None
    min_length, max_length = lengths
    with torch.no_grad():
        ids = model.generate(
            inputs['input_ids'],
            attention_mask=inputs.get('attention_mask'),
            num_beams=settings['num_beams'],
            min_length=min_length,
            max_length=max_length,
            length_penalty=settings['length_penalty'],
            no_repeat_ngram_size=3,
            early_stopping=True,
        )
    return _drop_incomplete_sentence(tokenizer.decode(ids[0], skip_special_tokens=True).strip())


def _drop_incomplete_sentence(text):
    """Remove a trailing sentence fragment left when generation hits max_length.

    Translation models hallucinate badly when given a fragment, so it is better
    to lose the fragment. Returns None if no complete sentence remains.
    """
    text = text.strip()
    if not text or re.search(r'[.!?\u0964]["\')\]]*$', text):
        return text or None
    ends = list(re.finditer(r'[.!?\u0964]["\')\]]*(?=\s)', text))
    return text[:ends[-1].end()] if ends else None


def _build_explanation(sentences, selected, scores, summary, mode):
    """Explain which sentences were selected and how strongly each scored."""
    def clip(s):
        s = s.strip()
        return s if len(s) <= 240 else s[:237].rsplit(' ', 1)[0] + '...'

    total = len(sentences)
    if mode == 'simple':
        points = '; '.join(clip(sentences[i]) for i in selected[:3])
        return f"The main points are: {points}\n\nShort summary:\n{summary}"

    if mode == 'detailed':
        lines = [f"{len(selected)} of {total} sentences were selected as the most representative "
                 "and least redundant content:"]
        for n, i in enumerate(selected, 1):
            lines.append(f"{n}. {clip(sentences[i])}")
            lines.append(f"   (sentence {i + 1} of {total}; relevance score {scores.get(i, 0):.2f})")
        lines.append('\nFull summary:')
        lines.append(summary)
        return '\n'.join(lines)

    ranked = sorted(selected, key=lambda i: scores.get(i, 0), reverse=True)[:5]
    lines = ['Key points (most relevant first):']
    lines += [f"{n}. {clip(sentences[i])}" for n, i in enumerate(ranked, 1)]
    lines.append('\nConcise summary:\n' + summary)
    return '\n'.join(lines)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def summarize(text, language='en', explain=False, summary_mode=None):
    """Summarize text and return the result in `language`.

    Returns a SummaryResult. Problems that did not prevent a summary (for
    example a missing translation model) are reported in `warnings`.
    """
    if isinstance(text, (list, tuple)):
        text = '\n\n'.join(t for t in text if t and isinstance(t, str))
    original_text = text or ''
    language = (language or 'en').strip().lower()
    mode = normalize_mode(summary_mode)
    result = SummaryResult(summary='', original_text=original_text, language=language)
    if not original_text.strip():
        return result

    source_lang = _detect_language(original_text)
    result.source_language = source_lang

    # Work in English: the summarizer and keyword extraction are English-only.
    # working_lang is the language of processing_text (and so of the raw summary).
    processing_text = original_text
    working_lang = source_lang
    n_source_sentences = 0
    if source_lang != 'en':
        source_sentences = _split_sentences(original_text)
        n_source_sentences = len(source_sentences)
        keep = _preselect(source_sentences, _PRESELECT_SENTENCES)
        to_translate = (original_text if len(keep) == n_source_sentences
                        else ' '.join(source_sentences[i] for i in keep))
        translated = _translate_text(to_translate, source_lang, 'en')
        if translated:
            processing_text = translated
            working_lang = 'en'
        else:
            result.warnings.append(
                f"Could not translate the input from '{source_lang}' to English; "
                "summarizing the original text, so quality may be lower.")

    sentences = _split_sentences(processing_text)
    if not sentences:
        return result

    # Size the summary by the whole document, even when only part was translated.
    settings = _mode_settings(mode, max(len(sentences), n_source_sentences))
    if language != 'en':
        # Translation compresses detail, so keep a little more source material.
        settings['top_k'] = min(len(sentences), max(settings['top_k'], int(settings['top_k'] * 1.5) + 1))

    embeddings = _compute_embeddings(sentences)
    if embeddings is not None:
        selected, scores = _hybrid_selector(embeddings, settings['top_k'])
    else:
        selected, scores = _frequency_selector(sentences, settings['top_k'])
    extractive_summary = ' '.join(sentences[i] for i in selected)

    summary = None
    if working_lang == 'en':
        try:
            summary = _abstractive_rewrite(extractive_summary, settings)
        except Exception as e:
            log.warning("Abstractive rewrite failed, using extractive summary: %s", e)
    if summary:
        result.method = 'hybrid'
    else:
        summary = extractive_summary
        result.method = 'extractive'

    explanation = _build_explanation(sentences, selected, scores, summary, mode) if explain else ''
    keywords = _extract_keywords(processing_text) if working_lang == 'en' else []

    if language != working_lang:
        translated_summary = _translate_text(summary, working_lang, language)
        if translated_summary:
            summary = translated_summary
            if explanation:
                explanation = _translate_text(explanation, working_lang, language) or explanation
        else:
            result.warnings.append(
                f"Translation to '{language}' is not available right now; showing the summary in "
                f"{'English' if working_lang == 'en' else working_lang}. Install the translation "
                "models (uv sync --extra ml) or enable online translation.")

    result.summary = summary
    result.explanation = explanation
    result.keywords = keywords
    return result


def funforfrontend(input_texts, language='en', explain=False, summary_mode='balanced'):
    """Backward-compatible wrapper returning (summary, original, explanation, keywords)."""
    r = summarize(input_texts, language=language, explain=explain, summary_mode=summary_mode)
    return r.summary, r.original_text, r.explanation or None, r.keywords
